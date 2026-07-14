#!/usr/bin/env python3
"""Parse the 3 Excel asset ledgers into JSON files (no DB access).

Extracts EVERY item row with all columns — including the Thừa/Thiếu
(surplus/missing) and Phẩm chất (quality) subcolumns that import_devices.py
ignores — plus a derived `loan` block interpreting the checkout/check-in
signals embedded in the sheets:

  Người sử dụng  a person name (thầy/cô/anh/chị/sếp …) = active borrower;
                 the column is also abused for free text (hư, Thiếu 1,
                 Nằm ở 504 …) which is kept as text, not a borrower.
  Thiếu          missing quantity = lost units.
  Ghi chú        "Đã trả" (returned), "… mượn" (borrowed by), "… giữ"
                 (held by), "yêu cầu/đề xuất trả" (return requested),
                 "phiếu điều chuyển …" (transfer slip to another room),
                 "Mã seri …" (serial number), "XIN GIẢM MÃ" (decommission).

Usage:
  python3 scripts/device-import/parse_ledgers.py            # ledgers in cwd
  python3 scripts/device-import/parse_ledgers.py --ledger-dir . --out-dir scripts/device-import/json

Writes one JSON per workbook (P402.json, P501.json, P502.json) plus a
combined all.json, and prints a per-signal summary. Requires openpyxl.
Row detection mirrors import_devices.py so device counts line up (648).
"""
import argparse, json, os, re, sys
from collections import Counter
from datetime import date, datetime

import openpyxl

SHEETS = {'Kiểm kê tài sản': 'tai_san', 'kiểm kê CCDC': 'ccdc'}

# Người sử dụng values that name a person (borrower), not free text.
PERSON_RE = re.compile(r'^(?:thầy|cô|anh|chị|sếp)\s+\S+', re.IGNORECASE)
SERIAL_RE = re.compile(r'mã seri\s*[:\s]*([\w./-]+)', re.IGNORECASE)
TRANSFER_RE = re.compile(r'điều chuyển.*?(\d{3})')
BORROWED_RE = re.compile(r'((?:thầy|cô|anh|chị|sếp|a|c)\.?\s+[^\s,;.]+)\s+(mượn|giữ)', re.IGNORECASE)


def norm(s):
    return re.sub(r'\s+', ' ', str(s).replace('\n', ' ')).strip() if s is not None else ''


def find_header(ws):
    for r in range(1, 20):
        if norm(ws.cell(r, 1).value).upper() == 'STT':
            return r
    sys.exit(f"error: no STT header row in sheet '{ws.title}'")


def map_cols(ws, hr):
    """Header-label mapping (layouts vary between files/sheets); the
    Thừa/Thiếu/quality labels live on the subheader row below."""
    roles = {}
    for c in range(1, ws.max_column + 1):
        lbl = norm(ws.cell(hr, c).value).lower()
        if not lbl:
            continue
        if 'mã' in lbl: roles.setdefault('code', c)
        elif lbl.startswith('tên'): roles.setdefault('name', c)
        elif lbl == 'đvt': roles.setdefault('unit', c)
        elif 'ngày' in lbl: roles.setdefault('date', c)
        elif 'sổ kế toán' in lbl: roles.setdefault('acct_qty', c)
        elif 'kiểm kê' in lbl and 'số lượng' in lbl: roles.setdefault('count_qty', c)
        elif 'người sử dụng' in lbl: roles.setdefault('user', c)
        elif 'hình ảnh' in lbl: roles.setdefault('image', c)
        elif 'vị trí' in lbl: roles.setdefault('location', c)
        elif 'ghi chú' in lbl: roles.setdefault('notes', c)
    for c in range(1, ws.max_column + 1):
        sub = norm(ws.cell(hr + 1, c).value).lower()
        if not sub:
            continue
        if 'thừa' in sub: roles.setdefault('surplus', c)
        elif 'thiếu' in sub: roles.setdefault('missing', c)
        elif 'còn tốt' in sub: roles.setdefault('q_good', c)
        elif 'kém' in sub: roles.setdefault('q_poor', c)
        elif 'mất' in sub: roles.setdefault('q_bad', c)
    return roles


def parse_date(v):
    if v in (None, ''):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})$', norm(v))
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d).strftime('%Y-%m-%d')
    except ValueError:
        return None


def as_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def loan_signals(user, image, notes):
    """Interpret the embedded checkout/check-in signals of one row."""
    text = f'{user} | {notes}'
    low = text.lower()
    loan = {
        'borrower': user if PERSON_RE.match(user) else None,
        'borrowerNote': None,        # "Sếp Đồng mượn" / "a Thắng giữ" in notes
        'returned': 'đã trả' in low,
        'returnRequested': ('yêu cầu trả' in low) or ('đề xuất trả' in low),
        'transferTo': None,          # phiếu điều chuyển → other room
        'decommissionRequested': 'xin giảm mã' in low,
    }
    m = BORROWED_RE.search(notes)
    if m:
        loan['borrowerNote'] = f'{m.group(1)} ({m.group(2)})'
    m = TRANSFER_RE.search(low)
    if m:
        loan['transferTo'] = f'P{m.group(1)}'
    return loan


def parse_workbook(path, room):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet, kind in SHEETS.items():
        ws = wb[sheet]
        hr = find_header(ws)
        roles = map_cols(ws, hr)
        cat = None
        for r in range(hr + 1, ws.max_row + 1):
            def g(role):
                c = roles.get(role)
                return ws.cell(r, c).value if c else None
            name, code = norm(g('name')), norm(g('code'))
            first = norm(ws.cell(r, 1).value)
            joined = ' '.join(norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1))
            if not name and not code:
                continue
            if 'tổng cộng' in joined.lower():
                continue
            if name and not code and not str(first).isdigit():
                cat = name
                continue
            if not code:
                continue
            user, image, notes = norm(g('user')), norm(g('image')), norm(g('notes'))
            serial = SERIAL_RE.search(notes)
            rows.append({
                'room': room,
                'sheet': kind,
                'excelRow': r,
                'category': cat or '',
                'code': code,
                'name': name,
                'unit': norm(g('unit')),
                'startDate': parse_date(g('date')),
                'bookQty': as_int(g('acct_qty')),
                'countedQty': as_int(g('count_qty')),
                'surplusQty': as_int(g('surplus')) or 0,
                'missingQty': as_int(g('missing')) or 0,
                # quality columns hold 'x' marks, not numbers → presence flags
                'quality': {
                    'good': bool(norm(g('q_good'))),
                    'poor': bool(norm(g('q_poor'))),
                    'bad': bool(norm(g('q_bad'))),
                },
                'user': user or None,
                'image': image or None,
                'position': norm(g('location')) or None,
                'notes': notes or None,
                'serialNumber': serial.group(1) if serial else None,
                'broken': 'hư' in (image.lower(), user.lower()),
                'loan': loan_signals(user, image, notes),
            })
    wb.close()
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--ledger-dir', default=os.environ.get('LEDGER_DIR', '.'))
    ap.add_argument('--out-dir', default=os.path.join(os.path.dirname(__file__), 'json'))
    args = ap.parse_args()

    files = {r: os.path.join(args.ledger_dir, f) for f in os.listdir(args.ledger_dir)
             if f.endswith('.xlsx') for r in ('P402', 'P501', 'P502') if r in f}
    if len(files) != 3:
        sys.exit(f"error: expected 3 ledgers in {args.ledger_dir}, found {sorted(files)}")

    os.makedirs(args.out_dir, exist_ok=True)
    all_rows = []
    for room, path in sorted(files.items()):
        rows = parse_workbook(path, room)
        out = os.path.join(args.out_dir, f'{room}.json')
        with open(out, 'w', encoding='utf-8') as fh:
            json.dump({'room': room, 'sourceFile': os.path.basename(path),
                       'rowCount': len(rows), 'rows': rows},
                      fh, ensure_ascii=False, indent=2)
        all_rows.extend(rows)
        print(f'{room}: {len(rows)} rows -> {out}')

    with open(os.path.join(args.out_dir, 'all.json'), 'w', encoding='utf-8') as fh:
        json.dump({'rowCount': len(all_rows), 'rows': all_rows},
                  fh, ensure_ascii=False, indent=2)

    c = Counter()
    for row in all_rows:
        loan = row['loan']
        c['total'] += 1
        c['borrower'] += bool(loan['borrower'])
        c['borrowerNote'] += bool(loan['borrowerNote'])
        c['missingQty>0'] += row['missingQty'] > 0
        c['returned'] += loan['returned']
        c['returnRequested'] += loan['returnRequested']
        c['transferTo'] += bool(loan['transferTo'])
        c['decommission'] += loan['decommissionRequested']
        c['serialNumber'] += bool(row['serialNumber'])
        c['broken'] += row['broken']
    print('\nsummary:')
    for k, v in c.items():
        print(f'  {k}: {v}')


if __name__ == '__main__':
    main()
