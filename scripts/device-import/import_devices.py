#!/usr/bin/env python3
"""Import devices from the 3 Excel ledgers directly into Supabase.

Maps straight from the .xlsx (deterministic) — no intermediate artifact.
Target is chosen entirely by env vars, so the same script serves local and
remote:

  # local
  SUPABASE_URL=http://127.0.0.1:54321 SUPABASE_KEY=<local secret> WIPE=1 \
    python3 import_devices.py

  # remote (prod)
  SUPABASE_URL=https://<ref>.supabase.co SUPABASE_KEY=<service_role> WIPE=1 \
    python3 import_devices.py

Requires: pip install supabase openpyxl
WIPE=1 deletes ALL existing devices first (clean re-import). Catalog rows
(Unknown manufacturer, groups) are upserted by name, so re-runs are safe.
Run import_photos.py afterwards to attach embedded photos.
"""
import argparse, os, re, sys
from collections import Counter, defaultdict
from datetime import date, datetime

import openpyxl
from supabase import create_client

UNIT_MAP = {'cái': 'piece', 'bộ': 'set', 'con': 'unit', 'thùng': 'box', 'chiếc': 'item'}
SHEETS = ('Kiểm kê tài sản', 'kiểm kê CCDC')


def env(name):
    v = os.environ.get(name)
    if not v:
        sys.exit(f"error: {name} env var is required (set it to choose the target).")
    return v


def norm(s):
    return re.sub(r'\s+', ' ', str(s).replace('\n', ' ')).strip() if s is not None else ''


def find_header(ws):
    for r in range(1, 20):
        if norm(ws.cell(r, 1).value).upper() == 'STT':
            return r


def map_cols(ws, hr):
    roles = {}
    for c in range(1, ws.max_column + 1):
        lbl = norm(ws.cell(hr, c).value).lower()
        if not lbl:
            continue
        if 'mã' in lbl: roles.setdefault('code', c)
        elif lbl.startswith('tên'): roles.setdefault('name', c)
        elif lbl == 'đvt': roles.setdefault('unit', c)
        elif 'ngày' in lbl: roles.setdefault('date', c)
        elif 'sổ kế toán' in lbl: roles.setdefault('acct_qty', c)
        elif 'kiểm kê' in lbl and 'số lượng' in lbl: roles.setdefault('count_qty', c)
        elif 'người sử dụng' in lbl: roles.setdefault('assignee', c)
        elif 'hình ảnh' in lbl: roles.setdefault('image', c)
        elif 'vị trí' in lbl: roles.setdefault('location', c)
        elif 'ghi chú' in lbl: roles.setdefault('notes', c)
    return roles


def parse_date(v):
    if v in (None, ''):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})$', norm(v))
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d).strftime('%Y-%m-%d')
    except ValueError:
        return None


def qty(acct, count):
    for v in (acct, count):
        try:
            n = int(float(v))
            if n >= 1:
                return n
        except (TypeError, ValueError):
            pass
    return 1


def notes_of(assignee, image, notes):
    parts = []
    if assignee:
        parts.append(f"Người sử dụng: {assignee}")
    if image and image.lower() == 'hư':
        parts.append("Tình trạng: hư")
    elif image:
        parts.append(image)
    if notes:
        parts.append(notes)
    return ' | '.join(parts) or None


def parse_ledgers(ledger_dir):
    files = {r: os.path.join(ledger_dir, f) for f in os.listdir(ledger_dir)
             if f.endswith('.xlsx') for r in ('P402', 'P501', 'P502') if r in f}
    if len(files) != 3:
        sys.exit(f"error: expected 3 ledgers in {ledger_dir}, found {sorted(files)}")
    rows = []
    for room, path in sorted(files.items()):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in SHEETS:
            ws = wb[sheet]
            hr = find_header(ws)
            roles = map_cols(ws, hr)
            cat = None
            for r in range(hr + 1, ws.max_row + 1):
                def g(role):
                    c = roles.get(role)
                    return ws.cell(r, c).value if c else None
                name, code = norm(g('name')), norm(g('code'))
                first = norm(ws.cell(r, 1).value)
                joined = ' '.join(norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1))
                if not name and not code:
                    continue
                if 'tổng cộng' in joined.lower():
                    continue
                if name and not code and not str(first).isdigit():
                    cat = name
                    continue
                if not code:
                    continue
                rows.append(dict(room=room, code=code, name=name, category=cat or '',
                                 unit=norm(g('unit')), acct_qty=g('acct_qty'),
                                 count_qty=g('count_qty'), date=parse_date(g('date')),
                                 assignee=norm(g('assignee')), location=norm(g('location')),
                                 notes=norm(g('notes')), image=norm(g('image'))))
        wb.close()
    return rows


def canonical_groups(rows):
    counts = Counter(r['category'] for r in rows if r['category'])
    ckey = lambda s: re.sub(r'[^a-z0-9]+', '', s.lower())
    by_key = defaultdict(list)
    for c in counts:
        by_key[ckey(c)].append(c)
    raw_to_canon = {v: max(vs, key=lambda x: counts[x]) for vs in by_key.values() for v in vs}
    return raw_to_canon, sorted(set(raw_to_canon.values()))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--ledger-dir', default=os.environ.get('LEDGER_DIR', '.'))
    args = ap.parse_args()
    url, key = env('SUPABASE_URL'), env('SUPABASE_KEY')
    wipe = os.environ.get('WIPE') == '1'

    rows = parse_ledgers(args.ledger_dir)
    raw_to_canon, groups = canonical_groups(rows)
    print(f"parsed {len(rows)} devices, {len(groups)} canonical groups")

    sk = create_client(url, key)
    sk.table('manufacturers').upsert({'name': 'Unknown'}, on_conflict='name').execute()
    mfr_id = sk.table('manufacturers').select('id').eq('name', 'Unknown').single().execute().data['id']
    for name in groups:
        sk.table('groups').upsert({'name': name}, on_conflict='name').execute()
    group_id = {g['name']: g['id']
                for g in sk.table('groups').select('id,name').execute().data}

    if wipe:
        sk.table('devices').delete().neq('id', '00000000-0000-0000-0000-000000000000').execute()
        print("wiped existing devices")

    payloads = []
    for r in rows:
        cat = raw_to_canon.get(r['category'])
        loc = r['room'] + (f" · {r['location']}" if r['location'] else '')
        payloads.append({
            'code': r['code'], 'name': r['name'] or '(no name)',
            'group_id': group_id[cat], 'manufacturer_id': mfr_id,
            'unit': UNIT_MAP.get(r['unit'].lower(), 'piece'),
            'quantity': qty(r['acct_qty'], r['count_qty']),
            'condition': 30 if r['image'].lower() == 'hư' else 100,
            'status': 'storage', 'location': loc, 'import_date': r['date'],
            'notes': notes_of(r['assignee'], r['image'], r['notes']),
        })
    for i in range(0, len(payloads), 100):
        sk.table('devices').insert(payloads[i:i + 100]).execute()

    total = sk.table('devices').select('id', count='exact').execute().count
    print(f"inserted {len(payloads)} devices; devices in DB now: {total}")


if __name__ == '__main__':
    main()
